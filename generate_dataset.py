#!/usr/bin/env python3
"""
BACKSTOP synthetic messy-data generator.
Creates a fictional nasi lemak stall dataset for testing extraction, reconciliation,
margin calculations, and source-row traceability.

Usage:
    python generate_dataset.py
    python generate_dataset.py --out data/demo_nasi_lemak
    python generate_dataset.py --seed 17

Dependency:
    python -m pip install openpyxl
"""
from __future__ import annotations

import argparse
import csv
import random
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError as exc:
    raise SystemExit(
        "Missing dependency 'openpyxl'. Install it with:\n"
        "  python -m pip install openpyxl"
    ) from exc


PRODUCTS = [
    {"name": "Nasi Lemak Biasa", "sell": 3.50, "rice": 0.42, "sambal": 0.28, "egg": 0.00, "pack": 0.18},
    {"name": "Nasi Lemak Telur", "sell": 4.50, "rice": 0.42, "sambal": 0.28, "egg": 0.55, "pack": 0.18},
    {"name": "Nasi Lemak Ayam", "sell": 7.50, "rice": 0.48, "sambal": 0.32, "egg": 0.55, "pack": 0.22, "chicken": 1.75},
    {"name": "Teh Ais", "sell": 2.00, "ice_tea": 0.48, "cup": 0.16},
]
START = date(2026, 9, 1)


def money(value: float) -> str:
    return f"RM {value:.2f}"


def write_receipts(root: Path, rng: random.Random) -> list[dict]:
    folder = root / "receipts_text"
    folder.mkdir(parents=True, exist_ok=True)
    rows = []
    suppliers = [
        ("Pasar Borong Seri Kembangan", "vegetables / rice"),
        ("Kedai Ayam Maju", "chicken"),
        ("Telur Segar Enterprise", "eggs"),
        ("Packaging Hub", "containers"),
        ("Mini Mart 88", "drinks / ice"),
    ]
    for i in range(1, 31):
        day = START + timedelta(days=(i - 1) * 1)
        supplier, category = suppliers[(i - 1) % len(suppliers)]
        amount = round(rng.uniform(38, 185), 2)
        payment = rng.choice(["DuitNow", "Cash", "Maybank QR", "cash"])
        # Deliberately inconsistent formatting, missing labels, and OCR-like noise.
        if i % 7 == 0:
            lines = [
                f"{supplier.upper()}",
                f"Date {day.strftime('%d/%m/%y')}",
                "ITEM      QTY      AMT",
                f"stock     {rng.randint(2, 8)}        {money(amount)}",
                f"TOTAL {money(amount)}",
                f"Paid via {payment}",
                f"Receipt No: R-{1000+i}",
            ]
        elif i % 5 == 0:
            lines = [
                f"{supplier} | {day.strftime('%Y-%m-%d')}",
                f"Ref: INV-{200+i}",
                f"Purchases ({category})",
                f"Grand total: RM{amount:.2f}",
                f"Payment={payment}",
                "Thank you",
            ]
        else:
            lines = [
                f"{supplier}",
                f"Receipt: R-{1000+i}",
                f"Date: {day.strftime('%d %b %Y')}",
                f"Category: {category}",
                f"Item: supplies",
                f"Amount: {money(amount)}",
                f"Payment method: {payment}",
            ]
        # A few intentionally ambiguous / imperfect records to test exception handling.
        if i in (9, 22):
            lines = [line for line in lines if not line.lower().startswith("date")]
            lines.append("date unclear")
        if i == 14:
            lines.append("TOTAL RM 118.80")  # conflicting duplicate total
        if i == 26:
            lines.append("Maybe personal item? ask owner")
        path = folder / f"receipt_{i:02d}.txt"
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        rows.append({
            "source_id": f"REC-{i:03d}",
            "source_file": f"receipts_text/{path.name}",
            "date": day.isoformat(),
            "supplier": supplier,
            "category": category,
            "amount_rm": amount,
            "payment_method": payment,
        })
    return rows


def write_whatsapp(root: Path, rng: random.Random) -> None:
    # Export resembles a copied WhatsApp chat, with informal Malay/English and corrections.
    lines = [
        "WhatsApp Chat Export - BACKSTOP demo (fictional)",
        "Chat: Nasi Lemak Kak Mira - Orders",
        "Exported: 2026-10-01",
        "",
    ]
    customers = ["Aina UPM", "Farid Kolej 14", "Jia Wen", "Amir", "Nadia", "Kumar", "Haziq"]
    for n in range(1, 46):
        day = START + timedelta(days=rng.randint(0, 29))
        hour = rng.randint(7, 11)
        minute = rng.choice([0, 5, 10, 15, 20, 30, 45])
        stamp = datetime(day.year, day.month, day.day, hour, minute)
        customer = rng.choice(customers)
        product = rng.choice(PRODUCTS)
        qty = rng.choice([1, 1, 1, 2, 3])
        total = product["sell"] * qty
        if n % 11 == 0:
            msg = f"{customer}: can reserve {qty} {product['name']}? bayar later ya"
        elif n % 9 == 0:
            msg = f"{customer}: {qty}x {product['name']} = RM{total:.0f} (rounding?); pickup 9am"
        elif n % 8 == 0:
            msg = f"{customer}: cancel order tadi sorry"
        else:
            msg = f"{customer}: order {qty} x {product['name']} RM {total:.2f}, pickup {hour+1}:{minute:02d}"
        lines.append(f"[{stamp.strftime('%d/%m/%Y, %I:%M %p')}] {msg}")
        if n % 13 == 0:
            correction_time = stamp + timedelta(minutes=2)
            lines.append(
                f"[{correction_time.strftime('%d/%m/%Y, %I:%M %p')}] {customer}: correction, make it {max(1, qty-1)} only"
            )
    lines.extend([
        "",
        "[01/10/2026, 09:12 AM] Mira: Some orders are paid cash, some QR. Don't count cancelled orders.",
        "[01/10/2026, 09:14 AM] Mira: Student promo last Friday was RM1 off ayam set, not sure if I wrote it down.",
        "[01/10/2026, 09:16 AM] Mira: Please separate confirmed sales from enquiries/reservations.",
    ])
    (root / "whatsapp_orders.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_broken_workbook(root: Path, rng: random.Random, receipts: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Sales"
    ws.append(["Date", "Product", "Qty Sold", "Unit Price (RM)", "Total", "Notes"])
    for n in range(1, 36):
        day = START + timedelta(days=rng.randint(0, 29))
        product = rng.choice(PRODUCTS)
        qty = rng.choice([8, 12, 15, 18, 20, 25, 30])
        unit = product["sell"]
        total = round(qty * unit, 2)
        # Deliberately broken spreadsheet: mixed dates, formulas, text numbers,
        # blank totals, duplicate rows, and a bad formula.
        date_value = day.strftime("%d/%m/%Y") if n % 4 == 0 else day.isoformat()
        qty_value = str(qty) if n % 6 == 0 else qty
        price_value = f"RM {unit:.2f}" if n % 7 == 0 else unit
        if n % 10 == 0:
            total_value = None
        elif n % 13 == 0:
            total_value = "=C999*D999"
        elif n % 8 == 0:
            total_value = f"{total:.2f} "
        else:
            total_value = total
        notes = rng.choice(["", "cash", "QR", "student promo?", "estimate", "check duplicate"])
        ws.append([date_value, product["name"], qty_value, price_value, total_value, notes])
        if n in (12, 27):
            # Duplicate an earlier record.
            ws.append([date_value, product["name"], qty_value, price_value, total_value, "possible duplicate copy"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243447")

    cost = wb.create_sheet("Costs")
    cost.append(["Item", "Cost", "Unit", "Updated?"])
    cost_rows = [
        ["Rice", 34.50, "5kg bag", "Aug?"],
        ["Eggs", "RM 14.00", "tray", "yes"],
        ["Chicken", 12.80, "kg", "old price"],
        ["Sambal ingredients", 22, "batch", ""],
        ["Containers", 0.18, "each", "estimate"],
        ["Rent", 450, "monthly", "fixed"],
        ["Gas", "unknown", "monthly", "ask Mira"],
        ["Helper allowance", 35, "day", "cash"],
        ["Electricity", 68.0, "monthly", "Sep bill"],
    ]
    for row in cost_rows:
        cost.append(row)
    for cell in cost[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")

    notes = wb.create_sheet("Read Me")
    notes.append(["Notes from owner (not fully reconciled)"])
    notes.append(["I think sales were around RM 180-260 most days"])
    notes.append(["Some personal groceries may be mixed into receipts"])
    notes.append(["Cash float not included"])
    notes.append(["Need to know whether delivery fees are included"])
    notes.append(["Don't treat estimates as confirmed"])
    out = root / "broken_sales_and_costs.xlsx"
    wb.save(out)


def write_manifest(root: Path, receipts: list[dict], seed: int) -> None:
    with (root / "receipt_manifest.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(receipts[0].keys()))
        writer.writeheader()
        writer.writerows(receipts)
    readme = f"""# BACKSTOP demo dataset

All records are synthetic and fictional. No real person or business is represented.
Generator seed: {seed}
Business scenario: small nasi lemak stall near a Malaysian university.

Files:
- `receipts_text/receipt_01.txt` ... `receipt_30.txt`: 30 receipt-like source documents.
- `receipt_manifest.csv`: generator ground truth for test/validation only; agents should not silently use it as evidence if simulating blind extraction.
- `whatsapp_orders.txt`: noisy chat export with cancellations, reservations, corrections, mixed language, and informal notes.
- `broken_sales_and_costs.xlsx`: intentionally inconsistent workbook with mixed date/number formats, missing totals, a bad formula, duplicates, and uncertain costs.

Important:
- This is demo data, not accounting advice or a representation of a real business.
- Do not treat the manifest as proof that the raw documents were correctly interpreted.
- Missing or contradictory records should be marked as unresolved rather than guessed.
"""
    (root / "DATASET_README.md").write_text(readme, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate BACKSTOP's synthetic messy dataset.")
    parser.add_argument("--out", default="data/demo_nasi_lemak", help="Output directory")
    parser.add_argument("--seed", type=int, default=42, help="Random seed for reproducibility")
    args = parser.parse_args()
    root = Path(args.out)
    root.mkdir(parents=True, exist_ok=True)
    rng = random.Random(args.seed)
    receipts = write_receipts(root, rng)
    write_whatsapp(root, rng)
    write_broken_workbook(root, rng, receipts)
    write_manifest(root, receipts, args.seed)
    print(f"BACKSTOP dataset created at: {root.resolve()}")
    print("Created 30 receipt text files, WhatsApp export, broken workbook, manifest, and README.")
    print("Next: inspect the raw files; use the manifest only as synthetic ground truth for validation.")


if __name__ == "__main__":
    main()
